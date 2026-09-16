"""Build the frozen-Chronos H16 comparison table from locked artifacts."""

from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from formal.config import MODEL_NAMES

BASELINE_MODELS = tuple(model for model in MODEL_NAMES if model != "smart_persistence")
TABLE_MODELS = BASELINE_MODELS[:1] + ("fm",) + BASELINE_MODELS[1:]
METRICS = ("rmse", "mae", "mbe", "r2")
METRIC_LABELS = {"rmse": "RMSE", "mae": "MAE", "mbe": "MBE", "r2": "R2"}
MODEL_LABELS = {
    "psrc": "PSRC",
    "fm": "Ours (PSRC+FM)",
    "timemixer": "TimeMixer",
    "ampdnet": "AMPDNet",
    "crossunet": "Cross-UNet",
    "timesnet": "TimesNet",
    "patchtst": "PatchTST",
    "itransformer": "iTransformer",
    "dlinear": "DLinear",
    "cyclenet": "CycleNet",
    "patchmlp": "PatchMLP",
    "tcn": "TCN",
    "frets": "FreTS",
    "lstm": "LSTM",
    "gru": "GRU",
    "timexer": "TimeXer",
    "persistence": "Persistence",
    "seasonal_naive": "Seasonal Naive",
}


def _metric(result: dict[str, object], name: str) -> float:
    key = "r2" if name == "r2" else f"{name}_physical"
    return float(result["test"][key])


def _result_index(results: Iterable[dict[str, object]]) -> dict[tuple[str, int, str, int], dict[str, object]]:
    index = {}
    for result in results:
        job = result["job"]
        key = (
            str(job["dataset"]),
            int(job["horizon"]),
            str(job["model"]),
            int(job["seed"]),
        )
        if key in index:
            raise ValueError(f"duplicate formal result: {key}")
        index[key] = result
    return index


def _load_fm_report(path: Path) -> dict[str, dict[str, object]]:
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if payload.get("protocol") != "single test unblinding; adapters loaded from locked development artefacts; alpha* fixed from validation":
        raise ValueError("FM report protocol is not the locked single-test-unblinding protocol")
    if payload.get("test_read_once") is not True:
        raise ValueError("FM report must document test_read_once=true")
    stations = payload.get("per_station")
    if not isinstance(stations, dict) or not stations:
        raise ValueError("FM report has no per_station records")
    return stations


def _validate_matrix(
    index: dict[tuple[str, int, str, int], dict[str, object]],
    datasets: tuple[str, ...],
    horizon: int,
    seed: int,
) -> None:
    for dataset in datasets:
        for model in BASELINE_MODELS:
            key = (dataset, horizon, model, seed)
            if key not in index:
                raise ValueError(f"missing formal result: {key}")
    unexpected = {
        (dataset, h, model, s)
        for dataset, h, model, s in index
        if dataset in datasets and h == horizon and s == seed and model not in MODEL_NAMES
    }
    if unexpected:
        raise ValueError(f"unknown model results: {sorted(unexpected)}")


def _best(values: dict[str, dict[str, float]], metric: str) -> str:
    candidates = ((model, row[metric]) for model, row in values.items())
    if metric == "r2":
        return max(candidates, key=lambda item: item[1])[0]
    if metric == "mbe":
        return min(candidates, key=lambda item: abs(item[1]))[0]
    return min(candidates, key=lambda item: item[1])[0]


def _style_sheet(sheet, last_column: int, last_row: int) -> None:
    navy = "1F4E78"
    green = "E2F0D9"
    red = "C00000"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C5"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color="1F2937")
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="666666")
    sheet["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 34
    for row in sheet.iter_rows(min_row=3, max_row=4, min_col=1, max_col=last_column):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=5, max_row=last_row, min_col=1, max_col=last_column):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="222222")
            cell.alignment = Alignment(horizontal="right" if cell.column > 2 else "left", vertical="center")
            if cell.column > 2:
                cell.number_format = "0.000000"
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 14
    for column in range(3, last_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 12
    # Red/bold rank 1 and underline rank 2, with MBE ranked by absolute bias.
    for row in range(5, last_row + 1):
        for metric_index, metric in enumerate(METRICS):
            entries = []
            for model_index, model in enumerate(TABLE_MODELS):
                cell = sheet.cell(row, 3 + model_index * 4 + metric_index)
                value = float(cell.value)
                entries.append((cell, value))
            entries.sort(key=(
                (lambda item: -item[1]) if metric == "r2"
                else (lambda item: abs(item[1])) if metric == "mbe"
                else (lambda item: item[1])
            ))
            for cell, _ in entries:
                cell.font = Font(name="Arial", size=10, color="222222")
                cell.fill = PatternFill(fill_type=None)
            entries[0][0].font = Font(name="Arial", size=10, bold=True, color=red)
            entries[0][0].fill = PatternFill("solid", fgColor=green)
            if len(entries) > 1:
                entries[1][0].font = Font(name="Arial", size=10, underline="single", color="222222")


def build_fm_comparison_workbook(
    results: list[dict[str, object]],
    fm_report_path: Path,
    output_path: Path,
) -> Path:
    """Build the four-station H16 table from baseline results plus locked FM test data."""
    index = _result_index(results)
    datasets = tuple(dict.fromkeys(str(result["job"]["dataset"]) for result in results))
    horizons = {int(result["job"]["horizon"]) for result in results}
    seeds = {int(result["job"]["seed"]) for result in results}
    if not datasets or horizons != {16} or seeds != {2026}:
        raise ValueError("FM table requires a non-empty H16, seed-2026 result set")
    required = ("dkasc_site31", "dkasc_site1a", "pvod_station02", "hkust")
    if datasets != required:
        raise ValueError(f"FM table requires datasets in this order: {list(required)}")
    _validate_matrix(index, datasets, 16, 2026)
    fm = _load_fm_report(fm_report_path)

    rows = []
    metadata = []
    for dataset in datasets:
        station = fm.get(dataset)
        if station is None:
            raise ValueError(f"FM report missing station: {dataset}")
        psrc_result = index[(dataset, 16, "psrc", 2026)]
        psrc_report_rmse = float(station["A_psrc"]["rmse_physical"])
        if not math.isclose(_metric(psrc_result, "rmse"), psrc_report_rmse, rel_tol=0.0, abs_tol=1e-5):
            raise ValueError(f"FM/PSRC alignment failed for {dataset}")
        values = {}
        for model in TABLE_MODELS:
            if model == "fm":
                source = station["FM_real"]
                values[model] = {metric: float(source[f"{metric}_physical"] if metric != "r2" else source[metric]) for metric in METRICS}
            else:
                result = index[(dataset, 16, model, 2026)]
                values[model] = {metric: _metric(result, metric) for metric in METRICS}
                metadata.append(_metadata_row(model, result))
        metadata.append(_fm_metadata_row(dataset, station))
        rows.append((dataset.upper(), "L96-H16", values))

    workbook = Workbook()
    comparison = workbook.active
    comparison.title = "Comparison"
    metadata_sheet = workbook.create_sheet("Run Metadata")
    params_sheet = workbook.create_sheet("PSRC Parameters")
    last_column = 2 + len(TABLE_MODELS) * 4
    comparison.append(["4-station H16 forecasting comparison with frozen Chronos correction"])
    comparison.append([
        "15 min resolution; L96 (24 h) to H16 (4 h); chronological 70%/15%/15%; seed 2026. Metrics use each station's native target scale; MBE = prediction - observation. Ours (PSRC+FM) uses validation-locked adapters and the single test unblinding on 2026-09-16, with no test-based reselection. Ours improves test RMSE on 3 of 4 stations; the real prior beats the shuffled-prior control on all 4."
    ])
    header = ["Station", "Setting"]
    metric_header = ["", ""]
    for model in TABLE_MODELS:
        header.extend([MODEL_LABELS[model], "", "", ""])
        metric_header.extend([METRIC_LABELS[metric] for metric in METRICS])
    comparison.append(header)
    comparison.append(metric_header)
    for model_index in range(len(TABLE_MODELS)):
        start = 3 + model_index * 4
        comparison.merge_cells(start_row=3, start_column=start, end_row=3, end_column=start + 3)
    for dataset, setting, values in rows:
        row = [dataset, setting]
        for model in TABLE_MODELS:
            row.extend(values[model][metric] for metric in METRICS)
        comparison.append(row)
    _style_sheet(comparison, last_column, comparison.max_row)

    metadata_headers = ["Model", "Dataset", "Setting", "Seed", "Parameter Source", "Epoch Budget", "Best Epoch", "Epochs Ran", "Early Stopped", "Stop Reason", "Patience", "Seconds"]
    metadata_sheet.append(metadata_headers)
    for row in metadata:
        metadata_sheet.append(row)
    _style_metadata(metadata_sheet, len(metadata_headers))

    # Keep the PSRC configuration already used by the selected result files.
    params_sheet.append(["Dataset", "Source", "Parameters"])
    for dataset in datasets:
        result = index[(dataset, 16, "psrc", 2026)]
        params_sheet.append([dataset.upper(), result["parameter_source"], json.dumps(result["parameters"], ensure_ascii=False, sort_keys=True)])
    _style_metadata(params_sheet, 3)
    params_sheet.column_dimensions["B"].width = 42
    params_sheet.column_dimensions["C"].width = 100

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _metadata_row(model: str, result: dict[str, object]) -> list[object]:
    return [
        MODEL_LABELS[model],
        str(result["job"]["dataset"]).upper(),
        f"L96-H{result['job']['horizon']}",
        int(result["job"]["seed"]),
        result.get("parameter_source", ""),
        result.get("epochs_requested"),
        result.get("best_epoch"),
        result.get("epochs_ran"),
        result.get("early_stopped"),
        result.get("stop_reason"),
        result.get("patience"),
        result.get("seconds"),
    ]


def _fm_metadata_row(dataset: str, station: dict[str, object]) -> list[object]:
    hp = station["locked_hyperparameters"]
    source = (
        f"Validation locked: prior={station['chosen_prior']}; source={station['locked_source']}; "
        f"lr={hp['lr']}; epsilon_c={hp['c']}; wd={hp['weight_decay']}; "
        f"huber_beta_std={hp['huber_beta_std']}; lambda_energy={hp['lambda_energy']}; "
        f"lambda_smooth={hp['lambda_smooth']}; alpha*={station['alpha_from_validation']}. "
        f"Single test unblinding; shuffled RMSE={station['FM_shuffled']['rmse_physical']}; no test reselection."
    )
    if not station.get("real_beats_psrc", True):
        source += " FM test result is worse than PSRC and is retained as observed."
    return ["Ours (PSRC+FM)", dataset.upper(), "L96-H16", 2026, source, 0, None, 0, False, "frozen_adapter", 0, None]


def _style_metadata(sheet, columns: int) -> None:
    navy = "1F4E78"
    for cell in sheet[1][:columns]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Arial", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, columns + 1):
        letter = sheet.cell(1, column).column_letter
        sheet.column_dimensions[letter].width = 22 if column != 5 else 88
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="222222")
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column == 5)
        sheet.row_dimensions[row[0].row].height = 48
