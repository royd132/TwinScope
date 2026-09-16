"""Aggregate formal task JSON files into the workbook payload."""

from __future__ import annotations

import json
import statistics
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from formal.config import (
    ALLOWED_HORIZONS,
    DATASET_NAMES,
    MODEL_NAMES,
    seeds_for_model,
)

MODELS = MODEL_NAMES
METRICS = ("rmse", "mae", "mbe", "r2")
# Rows display the seed mean together with the sample standard deviation.
METRIC_COLUMNS = tuple(
    (metric, statistic)
    for metric in METRICS
    for statistic in ("mean", "std")
)


def _seed_values(group: list[dict[str, object]], metric: str) -> list[float]:
    return [float(result["test"][f"{metric}_physical" if metric != "r2" else "r2"])
            for result in group]


def build_report_payload(results: list[dict[str, object]]) -> dict[str, object]:
    indexed: dict[tuple[str, int, str], list[dict[str, object]]] = {}
    for result in results:
        job = result["job"]
        key = (job["dataset"], int(job["horizon"]), job["model"])
        indexed.setdefault(key, []).append(result)

    # Campaign dimensions are inferred from the archived results so one
    # builder serves both the four-station H24 matrix and the five-station
    # H16 single-seed rerun without touching historical payloads.
    seen_datasets = {result["job"]["dataset"] for result in results}
    datasets = [name for name in DATASET_NAMES if name in seen_datasets]
    seen_horizons = {int(result["job"]["horizon"]) for result in results}
    horizons = [value for value in ALLOWED_HORIZONS if value in seen_horizons]
    experiment_seeds = sorted({int(result["job"]["seed"]) for result in results})
    anchor_seed = experiment_seeds[0]

    rows = []
    for dataset in datasets:
        for horizon in horizons:
            values = {}
            for model in MODELS:
                group = indexed.get((dataset, horizon, model), [])
                seeds = sorted(int(r["job"]["seed"]) for r in group)
                expected_seeds = [
                    seed
                    for seed in experiment_seeds
                    if seed in seeds_for_model(model)
                ]
                if seeds != expected_seeds:
                    raise ValueError(
                        f"incomplete seeds for {(dataset, horizon, model)}: "
                        f"{seeds}, expected {expected_seeds}"
                    )
                values[model] = {
                    metric: {
                        "mean": statistics.fmean(_seed_values(group, metric)),
                        "std": statistics.stdev(_seed_values(group, metric))
                        if len(group) > 1
                        else 0.0,
                    }
                    for metric in METRICS
                }
            best = {
                "rmse": min(MODELS, key=lambda m: values[m]["rmse"]["mean"]),
                "mae": min(MODELS, key=lambda m: values[m]["mae"]["mean"]),
                "mbe": min(MODELS, key=lambda m: abs(values[m]["mbe"]["mean"])),
                "r2": max(MODELS, key=lambda m: values[m]["r2"]["mean"]),
            }
            rows.append(
                {
                    "dataset": dataset.upper(),
                    "setting": f"L96-H{horizon}",
                    "forecast_hours": horizon / 4,
                    "values": values,
                    "best": best,
                }
            )

    metadata = []
    parameters = []
    for dataset in datasets:
        for horizon in horizons:
            for model in MODELS:
                for result in sorted(
                    indexed[(dataset, horizon, model)],
                    key=lambda r: int(r["job"]["seed"]),
                ):
                    metadata.append(
                        {
                            "model": model,
                            "dataset": dataset.upper(),
                            "setting": f"L96-H{horizon}",
                            "seed": result["job"]["seed"],
                            "parameter_source": result["parameter_source"],
                            "epoch_budget": result["epochs_requested"],
                            "best_epoch": result["best_epoch"],
                            "epochs_ran": result["epochs_ran"],
                            "early_stopped": result["early_stopped"],
                            "stop_reason": result["stop_reason"],
                            "patience": result["patience"],
                            "seconds": result["seconds"],
                        }
                    )
                anchor = next(
                    r
                    for r in indexed[(dataset, horizon, model)]
                    if int(r["job"]["seed"]) == anchor_seed
                )
                if model == "psrc":
                    parameters.append(
                        {
                            "dataset": dataset.upper(),
                            "source": anchor["parameter_source"],
                            "params": anchor["parameters"],
                        }
                    )
    seed_text = ", ".join(map(str, experiment_seeds))
    horizon_text = "/".join(f"H{value}" for value in horizons)
    hours_text = "/".join(f"{value / 4:g}" for value in horizons)
    return {
        "title": f"{len(datasets)}-station {horizon_text} forecasting comparison",
        "note": (
            f"15 min resolution; L96 (24 h) to {horizon_text} ({hours_text} h); "
            f"seeds {seed_text}; each cell is the seed mean with the sample "
            "standard deviation beside it. RMSE, MAE and MBE use each station's "
            "native target scale; MBE = prediction - observation."
        ),
        "models": list(MODELS),
        "metrics": list(METRICS),
        "metric_columns": list(METRIC_COLUMNS),
        "seeds": experiment_seeds,
        "rows": rows,
        "metadata": metadata,
        "parameters": parameters,
    }


def write_report_payload(results: list[dict[str, object]], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(build_report_payload(results), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return path


_MODEL_LABELS = {
    "psrc": "PSRC",
    "timemixer": "TimeMixer",
    "ampdnet": "AMPDNet",
    "crossunet": "Cross-UNet",
    "timesnet": "TimesNet",
    "patchtst": "PatchTST",
    "itransformer": "iTransformer",
    "dlinear": "DLinear",
    "cyclenet": "CycleNet",
    "patchmlp": "PatchMLP",
    "tcn": "TCN",
    "frets": "FreTS",
    "lstm": "LSTM",
    "gru": "GRU",
    "timexer": "TimeXer",
    "persistence": "Persistence",
    "seasonal_naive": "Seasonal Naive",
    "smart_persistence": "Smart Persistence",
}
_METRIC_LABELS = {
    ("rmse", "mean"): "RMSE", ("rmse", "std"): "RMSE SD",
    ("mae", "mean"): "MAE", ("mae", "std"): "MAE SD",
    ("mbe", "mean"): "MBE", ("mbe", "std"): "MBE SD",
    ("r2", "mean"): "R2", ("r2", "std"): "R2 SD",
}


def build_workbook(payload_path: Path, output_path: Path) -> None:
    payload = json.loads(Path(payload_path).read_text(encoding="utf-8"))
    models = payload["models"]
    metric_columns = [tuple(column) for column in payload["metric_columns"]]
    group_width = len(metric_columns)
    last_column = 2 + len(models) * group_width

    workbook = Workbook()
    comparison = workbook.active
    comparison.title = "Comparison"
    metadata = workbook.create_sheet("Run Metadata")
    parameters = workbook.create_sheet("PSRC Parameters")
    navy = "1F4E78"
    green = "E2F0D9"

    comparison.append([payload["title"]])
    comparison.append([payload["note"]])
    header_model = ["Station", "Setting"]
    header_metric = ["", ""]
    for model in models:
        header_model.extend([_MODEL_LABELS[model]] + [""] * (group_width - 1))
        header_metric.extend(_METRIC_LABELS[column] for column in metric_columns)
    comparison.append(header_model)
    comparison.append(header_metric)
    for column in range(3, last_column + 1, group_width):
        comparison.merge_cells(
            start_row=3, start_column=column,
            end_row=3, end_column=column + group_width - 1,
        )

    for row in payload["rows"]:
        values = [row["dataset"], row["setting"]]
        for model in models:
            for metric, statistic in metric_columns:
                values.append(row["values"][model][metric][statistic])
        comparison.append(values)
        excel_row = comparison.max_row
        for model_index, model in enumerate(models):
            for metric_index, (metric, statistic) in enumerate(metric_columns):
                cell = comparison.cell(
                    excel_row, 3 + model_index * group_width + metric_index
                )
                cell.number_format = "0.000000"
                if statistic == "mean" and row["best"][metric] == model:
                    cell.fill = PatternFill("solid", fgColor=green)
                    cell.font = Font(bold=True)
    comparison.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    comparison.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    comparison["A1"].font = Font(size=15, bold=True)
    comparison["A2"].font = Font(size=9, italic=True, color="595959")
    for row in comparison.iter_rows(min_row=3, max_row=4, min_col=1, max_col=last_column):
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
    comparison.freeze_panes = "C5"
    comparison.column_dimensions["A"].width = 20
    comparison.column_dimensions["B"].width = 14
    for column in range(3, last_column + 1):
        comparison.column_dimensions[get_column_letter(column)].width = 11

    metadata_headers = [
        "Model",
        "Dataset",
        "Setting",
        "Seed",
        "Parameter Source",
        "Epoch Budget",
        "Best Epoch",
        "Epochs Ran",
        "Early Stopped",
        "Stop Reason",
        "Patience",
        "Seconds",
    ]
    metadata.append(metadata_headers)
    for row in payload["metadata"]:
        metadata.append(
            [
                _MODEL_LABELS[row["model"]],
                row["dataset"],
                row["setting"],
                row["seed"],
                row["parameter_source"],
                row["epoch_budget"],
                row["best_epoch"],
                row["epochs_ran"],
                row["early_stopped"],
                row["stop_reason"],
                row["patience"],
                row["seconds"],
            ]
        )
    _style_table(metadata, len(metadata_headers), navy)

    parameter_headers = [
        "Dataset",
        "Source",
        "d_model",
        "Heads",
        "Patch hours",
        "GTR period",
        "Dropout",
        "Learning rate",
        "Weight decay",
        "Batch",
        "Optimizer",
        "Scheduler",
        "Base loss",
        "Huber delta",
        "Sampling",
        "Daylight weight",
        "Ramp weight",
        "Max gate",
        "Max correction",
        "Semantic weight",
        "Residual alpha",
        "Base anchor",
    ]
    parameters.append(parameter_headers)
    for row in payload["parameters"]:
        p = row["params"]
        parameters.append(
            [
                row["dataset"],
                row["source"],
                p["d_model"],
                p["heads"],
                " / ".join(map(str, p["patch_hours"])),
                p["gtr_period"],
                p["dropout"],
                p["lr"],
                p["weight_decay"],
                p["batch"],
                p["optimizer"],
                p["scheduler"],
                p["loss_base"],
                p["huber_delta"],
                p["sampling_policy"],
                p["daylight_weight"],
                p["ramp_weight"],
                p["max_gate"],
                p["max_correction"],
                p["semantic_loss_weight"],
                p["residual_alpha"],
                p["base_anchor_weight"],
            ]
        )
    _style_table(parameters, len(parameter_headers), navy)
    parameters.column_dimensions["B"].width = 38
    parameters.column_dimensions["E"].width = 18
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _style_table(sheet, columns: int, header_color: str) -> None:
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, columns + 1):
        letter = sheet.cell(1, column).column_letter
        values = [
            str(sheet.cell(row, column).value or "")
            for row in range(1, sheet.max_row + 1)
        ]
        sheet.column_dimensions[letter].width = min(
            28, max(11, max(map(len, values)) + 2)
        )
