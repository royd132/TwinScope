import pytest
from openpyxl import load_workbook

from formal.config import MODEL_NAMES, SEEDS, seeds_for_model
from formal.report import (
    build_report_payload,
    build_workbook,
    write_report_payload,
)

DATASETS = ("dkasc_site31", "dkasc_site9a", "pvod_station00", "pvod_station01")
HORIZONS = (24,)

PSRC_PARAMS = {
    "d_model": 64,
    "heads": 4,
    "patch_hours": [1, 2, 4, 8],
    "gtr_period": 24,
    "dropout": 0.1,
    "lr": 0.001,
    "weight_decay": 0.0001,
    "batch": 128,
    "optimizer": "adam",
    "scheduler": "none",
    "loss_base": "mse",
    "huber_delta": 1.0,
    "sampling_policy": "sequential",
    "daylight_weight": 0.25,
    "ramp_weight": 0.0,
    "max_gate": 0.5,
    "max_correction": 0.25,
    "semantic_loss_weight": 0.2,
    "residual_alpha": 1.0,
    "base_anchor_weight": 1.0,
}


def _result(dataset, horizon, model, seed, value):
    return {
        "status": "complete",
        "config_fingerprint": f"fingerprint-{dataset}-{horizon}-{model}-{seed}",
        "headline_test_evaluation_passes": 1,
        "job": {
            "key": f"{dataset}-H{horizon}-{model}-{seed}",
            "dataset": dataset,
            "model": model,
            "seq_len": 96,
            "horizon": horizon,
            "seed": seed,
        },
        "test": {
            "rmse_physical": value,
            "mae_physical": value / 2,
            "mbe_physical": (-0.5 if model == "timemixer" else 1) * value / 10,
            "r2": 1 - value / 10,
        },
        "parameter_source": "tuned" if model == "psrc" else "baseline_fixed",
        "parameters": dict(PSRC_PARAMS) if model == "psrc" else {},
        "epochs_requested": 100,
        "epochs_ran": 20 if model not in ("persistence", "seasonal_naive", "smart_persistence") else 0,
        "best_epoch": 10,
        "early_stopped": True,
        "stop_reason": "patience",
        "patience": 10,
        "seconds": 12.5,
    }


@pytest.fixture
def fake_results():
    model_value = {
        "psrc": 1.0,
        "timemixer": 1.1,
        "ampdnet": 1.2,
        "crossunet": 1.3,
        "timesnet": 1.4,
        "patchtst": 1.5,
        "itransformer": 1.6,
        "dlinear": 1.7,
        "cyclenet": 1.8,
        "patchmlp": 1.9,
        "tcn": 2.3,
        "frets": 2.4,
        "lstm": 2.5,
        "gru": 2.6,
        "timexer": 2.7,
        "persistence": 2.0,
        "seasonal_naive": 2.1,
        "smart_persistence": 2.2,
    }
    results = []
    for dataset in DATASETS:
        for horizon in HORIZONS:
            for model in MODEL_NAMES:
                for seed_index, seed in enumerate(seeds_for_model(model)):
                    results.append(
                        _result(
                            dataset,
                            horizon,
                            model,
                            seed,
                            model_value[model] + seed_index * 0.01,
                        )
                    )
    return results


def test_report_payload_has_four_rows_and_thirteen_model_groups(fake_results):
    payload = build_report_payload(fake_results)

    assert len(payload["rows"]) == 4
    assert payload["models"] == list(MODEL_NAMES)
    assert payload["metrics"] == ["rmse", "mae", "mbe", "r2"]
    assert payload["seeds"] == list(SEEDS)
    assert len(payload["metadata"]) == 4 * 13 * 3 + 4 * 5
    assert len(payload["parameters"]) == 4


def test_report_aggregates_seed_mean_and_std(fake_results):
    row = build_report_payload(fake_results)["rows"][0]

    assert row["values"]["psrc"]["rmse"]["mean"] == pytest.approx(1.01)
    assert row["values"]["psrc"]["rmse"]["std"] == pytest.approx(0.01)
    assert row["values"]["psrc"]["mae"]["mean"] == pytest.approx(0.505)
    assert row["values"]["psrc"]["r2"]["mean"] == pytest.approx(0.899)


def test_report_marks_metric_specific_best_values(fake_results):
    row = build_report_payload(fake_results)["rows"][0]

    assert row["best"] == {
        "rmse": "psrc",
        "mae": "psrc",
        "mbe": "timemixer",
        "r2": "psrc",
    }


def test_report_rejects_incomplete_seed_matrix(fake_results):
    with pytest.raises(ValueError, match="incomplete seeds"):
        build_report_payload(fake_results[:-1])


def test_artifact_builder_creates_expected_workbook(tmp_path, fake_results):
    payload_path = write_report_payload(fake_results, tmp_path / "report.json")
    workbook_path = tmp_path / "comparison.xlsx"

    build_workbook(payload_path, workbook_path)

    assert workbook_path.is_file()
    workbook = load_workbook(workbook_path, data_only=False)
    assert workbook.sheetnames == ["Comparison", "Run Metadata", "PSRC Parameters"]
    comparison = workbook["Comparison"]
    assert comparison.max_row == 8
    assert comparison.max_column == 2 + 18 * 8
    assert comparison["A5"].value == "DKASC_SITE31"
    assert comparison["B5"].value == "L96-H24"
    # First PSRC RMSE cell holds the seed mean and is the highlighted best.
    assert comparison["C5"].value == pytest.approx(1.01)
    assert comparison["C5"].font.bold is True
    assert comparison["D5"].value == pytest.approx(0.01)
    assert len(workbook["Run Metadata"]["A"]) == 1 + 4 * 13 * 3 + 4 * 5
    assert len(workbook["PSRC Parameters"]["A"]) == 5
